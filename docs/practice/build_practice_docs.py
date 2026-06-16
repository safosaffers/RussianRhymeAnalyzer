#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Оформляет каждую из 4 частей практики (01..04_*.md) отдельным Word-документом
# по ГОСТ. Переиспользует стили из build_report.py (TNR 14, интервал 1.5, поля,
# заголовки, нумерация страниц). Запуск: python docs/practice/build_practice_docs.py
import os
import re

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

import build_report as R   # setup_styles, page_footer, h1, h2, para, bullets, table

HERE = os.path.dirname(os.path.abspath(__file__))

DOCS = [
    ("01_sources.md",      "Задание 1. Работа с источниками",
     "Обзор литературы, аналитический обзор (SOTA) и библиография"),
    ("02_plan_gantt.md",   "Задание 2. План-график работы над ВКР",
     "Помесячный план на 2 года, диаграмма Ганта и контрольные точки"),
    ("03_data.md",         "Задание 3. Работа с данными",
     "Поиск, описание наборов данных, предобработка и протокол сбора"),
    ("04_specification.md", "Задание 4. Спецификация системы",
     "Математическая постановка, метрики, окружение и пайплайн данных"),
]

# доп. раздел для спецификации — как мы работаем с ИИ (оценка — коррекция)
AI_SECTION = """## 5. Взаимодействие с ИИ: оценка — коррекция
Хотя тема ВКР относится к искусственному интеллекту, ядро работы — математика
оценки рифмы; именно метрики направляют исправление текста. Взаимодействие с
языковой моделью построено как цикл с обратной связью.
- Генерация: модель выдаёт N вариантов стиха на тему (best-of-N).
- Оценка: детектор рифмы численно оценивает каждый вариант (RPST или Юкава).
- Разметка: по результату оценки помечаются окончания строк — какие рифмуются (оставить), а какие нет (переписать).
- Коррекция: модели передаётся этот разбор и запрос переписать только плохо рифмующиеся строки, сохранив удачные рифмы, и проверить читаемость (естественный порядок слов) и форму (размер, число строк, схема).
Таким образом языковая модель выступает инструментом, а решения о том, что
исправлять, принимаются на основе численных метрик качества рифмы.
"""


def _clean(s: str) -> str:
    s = re.sub(r'!?\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', s)   # ссылки
    s = re.sub(r'\*\*([^*]+)\*\*', r'\1', s)                  # жирный
    s = re.sub(r'`([^`]+)`', r'\1', s)                       # код
    return s.strip()


def _row(line: str):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def _is_special(s: str) -> bool:
    return (s.startswith('#') or s.startswith('>') or s.startswith('|')
            or s.startswith('```') or bool(re.match(r'^[-*]\s', s))
            or bool(re.match(r'^\d+\.\s', s)))


def md_to_doc(doc, md: str):
    lines = md.splitlines()
    N, i, first_h1 = len(lines), 0, True
    while i < N:
        s = lines[i].strip()
        if not s:
            i += 1; continue

        if s.startswith('### '):
            R.h2(doc, _clean(s[4:])); i += 1; continue
        if s.startswith('## '):
            R.h1(doc, _clean(s[3:]), page_break=False, toc=False); i += 1; continue
        if s.startswith('# '):
            if first_h1:        # заголовок документа уже на титуле — пропускаем
                first_h1 = False
            else:
                R.h1(doc, _clean(s[2:]), page_break=False, toc=False)
            i += 1; continue

        if s.startswith('```'):                      # код/диаграмма
            i += 1
            while i < N and not lines[i].lstrip().startswith('```'):
                i += 1
            i += 1
            R.para(doc, "(Диаграмма Ганта (mermaid) — см. файл 02_plan_gantt.md.)"
                   ).runs[0].italic = True
            continue

        if s.startswith('|') and i + 1 < N and set(
                lines[i + 1].replace('|', '').replace(':', '').strip()) <= {'-', ' '}:
            header = _row(s); i += 2; rows = []
            while i < N and lines[i].strip().startswith('|'):
                rows.append([_clean(c) for c in _row(lines[i])]); i += 1
            R.table(doc, header, rows); continue

        if s.startswith('>'):                        # цитата (склеиваем строки)
            buf = []
            while i < N and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip()); i += 1
            p = R.para(doc, _clean(' '.join(buf)))
            for r in p.runs:
                r.italic = True
            continue

        if re.match(r'^[-*]\s', s):                  # маркированный список
            items = []
            while i < N:
                ls = lines[i].strip()
                if re.match(r'^[-*]\s', ls):
                    items.append(_clean(ls[2:])); i += 1
                elif ls and not _is_special(ls):     # продолжение пункта (перенос)
                    items[-1] += ' ' + _clean(ls); i += 1
                else:
                    break
            R.bullets(doc, items); continue

        if re.match(r'^\d+\.\s', s):                  # нумерованный список
            items = []
            while i < N:
                ls = lines[i].strip()
                if re.match(r'^\d+\.\s', ls):
                    items.append(_clean(re.sub(r'^\d+\.\s*', '', ls))); i += 1
                elif ls and not _is_special(ls):
                    items[-1] += ' ' + _clean(ls); i += 1
                else:
                    break
            for it in items:
                p = doc.add_paragraph(style="List Number"); p.add_run(it)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            continue

        buf = [s]; i += 1                            # абзац (склеиваем переносы)
        while i < N and lines[i].strip() and not _is_special(lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        R.para(doc, _clean(' '.join(buf)))


def title_page(doc, task_title, subtitle):
    def line(t, bold=False, size=14, before=0, after=0):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)

    line("Министерство науки и высшего образования Российской Федерации")
    line("Федеральное государственное бюджетное образовательное учреждение")
    line("высшего образования")
    line("«Новгородский государственный университет имени Ярослава Мудрого»", bold=True)
    line("Политехнический институт")
    line("Кафедра информационных технологий и систем", after=30)

    line("ОТЧЁТ по учебной практике", bold=True, size=16, before=24)
    line(task_title, bold=True, size=15, before=10, after=6)
    line(subtitle, size=13, after=20)
    line("по направлению 09.04.01 «Информатика и вычислительная техника»,")
    line("профиль «Искусственный интеллект»", after=40)
    line("Выполнил: студент группы 5095  С. Каримов", after=4)
    line("Руководитель: ______________ / ____________________", after=40)
    line("Великий Новгород — 2026", before=30)
    doc.add_paragraph().paragraph_format.page_break_before = True


def build():
    out = []
    for fn, task_title, subtitle in DOCS:
        md = open(os.path.join(HERE, fn), encoding="utf-8").read()
        if fn.startswith("04"):
            md = md.rstrip() + "\n\n" + AI_SECTION
        doc = Document()
        R.setup_styles(doc)
        R.page_footer(doc)
        title_page(doc, task_title, subtitle)
        md_to_doc(doc, md)
        name = {"01": "Практика_1_Источники.docx",
                "02": "Практика_2_План-график.docx",
                "03": "Практика_3_Данные.docx",
                "04": "Практика_4_Спецификация.docx"}[fn[:2]]
        path = os.path.join(HERE, name)
        doc.save(path)
        out.append(name)
    print("Готово:", ", ".join(out))


if __name__ == "__main__":
    build()
