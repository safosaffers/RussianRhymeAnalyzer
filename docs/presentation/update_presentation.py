#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Добавляет в Rhymer_presentation_final.pptx (НЕ затирая правки):
#  - слайд «Как мы работаем с ИИ» (оценка → коррекция);
#  - вместо старого демо-слайда с 3 картинками — по слайду на каждый скриншот.
# Новые слайды строятся с нуля (фон-картинка + заголовок + контент), без
# клонирования. Делает бэкап. Запуск: python docs/presentation/update_presentation.py
import os
import shutil

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
PPTX = os.path.join(HERE, "Rhymer_presentation_final.pptx")
BG = os.path.join(ROOT, "src", "View", "assets", "bg.png")
SHOTS = os.path.join(HERE, "shots")

NAVY = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
YELLOW = RGBColor(0xFB, 0xBF, 0x24)


def add_bg(prs, slide):
    slide.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)


def add_title(slide, text, w):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Emu(int(w * 0.06)), Emu(int(w * 0.025)),
                                 Emu(int(w * 0.88)), Pt(70))
    box.fill.solid(); box.fill.fore_color.rgb = NAVY
    box.line.fill.background()
    tf = box.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = WHITE
    return box


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_bg(prs, slide)
    return slide


def ai_loop_slide(prs):
    W, H = prs.slide_width, prs.slide_height
    slide = new_slide(prs)
    add_title(slide, "Как мы работаем с ИИ: оценка → коррекция", W)
    steps = [
        ("1", "Генерация", "ИИ пишет N вариантов стиха (best-of-N)."),
        ("2", "Оценка", "Наши метрики численно оценивают рифму (RPST / Юкава)."),
        ("3", "Разметка", "Метим окончания: где рифма есть — оставить, где нет — переписать."),
        ("4", "Коррекция", "Просим ИИ перерифмовать плохое, сохранив хорошее, и проверить читаемость и форму."),
    ]
    n = len(steps)
    gap = int(W * 0.02)
    cw = int((W * 0.88 - gap * (n - 1)) / n)
    x0 = int(W * 0.06)
    top = Pt(150)
    ch = int(H * 0.55)
    for i, (num, head, desc) in enumerate(steps):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Emu(x0 + i * (cw + gap)), Emu(int(top)),
                                      Emu(cw), Emu(ch))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = NAVY; card.line.width = Pt(2.5)
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(12)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.LEFT
        r = p0.add_run(); r.text = num
        r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = RGBColor(0xEC, 0x48, 0x99)
        p1 = tf.add_paragraph(); rr = p1.add_run(); rr.text = head
        rr.font.size = Pt(22); rr.font.bold = True; rr.font.color.rgb = NAVY
        p2 = tf.add_paragraph(); rd = p2.add_run(); rd.text = desc
        rd.font.size = Pt(16); rd.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    # подпись-вывод снизу
    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Emu(int(W * 0.06)), Emu(int(H * 0.82)),
                                  Emu(int(W * 0.88)), Pt(54))
    note.fill.solid(); note.fill.fore_color.rgb = NAVY; note.line.fill.background()
    tf = note.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "ИИ — инструмент; ядро работы — математика оценки рифмы, которая и направляет правки."
    r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = YELLOW
    return slide


def shot_slide(prs, title, img):
    W, H = prs.slide_width, prs.slide_height
    slide = new_slide(prs)
    add_title(slide, title, W)
    iw, ih = Image.open(img).size
    maxw, maxh = int(W * 0.9), int(H * 0.74)
    ar = iw / ih
    if maxw / maxh > ar:
        h = maxh; w = int(h * ar)
    else:
        w = maxw; h = int(w / ar)
    left = int((W - w) / 2); top = int(H * 0.18)
    slide.shapes.add_picture(img, Emu(left), Emu(top), width=Emu(w), height=Emu(h))
    return slide


def reorder(prs, order):
    lst = prs.slides._sldIdLst
    elems = list(lst)
    for e in elems:
        lst.remove(e)
    for i in order:
        lst.append(elems[i])


def main():
    shutil.copy2(PPTX, os.path.join(HERE, "Rhymer_presentation_final_BACKUP.pptx"))
    prs = Presentation(PPTX)
    n0 = len(prs.slides._sldIdLst)             # исходное число слайдов (15)

    ai_loop_slide(prs)                          # индекс n0
    shots = [
        ("Демонстрация: генерация стиха", "01_generation.png"),
        ("Демонстрация: раскраска рифм", "02_coloring.png"),
        ("Демонстрация: тёмная тема", "03_dark.png"),
        ("Демонстрация: настройки (провайдер и ключ)", "04_settings.png"),
    ]
    for title, fn in shots:                     # индексы n0+1 .. n0+4
        shot_slide(prs, title, os.path.join(SHOTS, fn))

    ai = n0
    s1, s2, s3, s4 = n0 + 1, n0 + 2, n0 + 3, n0 + 4
    DEMO = 12                                    # старый демо-слайд с 3 картинками — выкинуть
    # 0..10, AI, 11(Итог?), скриншоты, 13(Результаты), 14(Спасибо)
    order = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, ai, 11, s1, s2, s3, s4, 13, 14]
    assert DEMO not in order
    reorder(prs, order)

    prs.save(PPTX)
    print("Сохранено:", PPTX, "| слайдов:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
